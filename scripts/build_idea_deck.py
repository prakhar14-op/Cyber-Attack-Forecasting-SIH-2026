"""Render `docs/idea_round.md` into the 6-slide IDEA-ROUND .pptx.

    C:/sih26/.venv/Scripts/python.exe scripts/build_idea_deck.py

WHICH DECK THIS IS, AND WHICH IT IS NOT
---------------------------------------
SIH runs two evaluations and this repository now has a deck for each.

  * `scripts/build_deck.py`  -> `docs/deck/sih26153_technical_deck.pptx`
    FIVE slides, the PS's deliverable-10 technical presentation, for the December
    Grand Finale. It leads with limitations because finale judges open the repo.

  * `scripts/build_idea_deck.py` (this file) -> `docs/deck/sih26153_idea_deck.pptx`
    SIX slides, the September idea-round screening submission, on the official
    SIH IDEA template's fixed section list. Screening judges typically never open
    the repository, read six slides in about ninety seconds, and are comparing
    against decks that claim near-perfect accuracy with no caveats.

This script never touches the technical deck and the technical deck's builder
never touches this one. They share machinery, not content: everything imported
from `build_deck` below is parsing, layout and the two banned-claim families that
apply to any judge-facing surface this project produces. `build_deck.py` is NOT
modified — the shared pieces are imported from it, which means the idea deck
inherits a fix to the markdown parser or the chevron layout for free, and also
means a future extraction of those pieces into a module is a pure move.

WHY A BUILDER AND NOT A HAND-MADE .pptx
---------------------------------------
The same reason `build_deck.py` gives, and it bites harder here. This project's
named failure mode is CLAIMS THAT OUTRUN THE CODE, and the idea round is the
surface with the least friction against it: nobody in that room can check. So the
.pptx is generated, the content lives in reviewable markdown, and the numbers are
not in the markdown either.

THE OFFICIAL FORMAT, AND WHERE IT CAME FROM
-------------------------------------------
`SECTIONS` is the official SIH IDEA presentation format: six slides, fixed
titles, submitted as PDF. It was read off the official
`SIH2024_IDEA_Presentation_Format.pdf` slide by slide and cross-checked against
the SIH2026 revision of the same template, which keeps the section list and the
six-slide/PDF rule. `docs/idea_round.md` carries the table and the provenance.

The build fails if the document renames, reorders, drops or adds a section. A
screening judge scores against that list; a deck that quietly drops "Feasibility
and Viability" loses those marks whatever is on the slide instead.

NO RESULTS NUMBER IS TYPED ANYWHERE
-----------------------------------
`docs/architecture.md` §4's results table is being regenerated under a new
anonymisation key, and a `net24_bucket` ablation may move the F1/recall columns.
A number pasted into the markdown today ships stale in September. So:

    {{measured:MODEL:COLUMN}}   that model's cell in the §4 table
    {{lead_min:MODEL}}          that model's median-lead cell, in whole minutes
    {{clopper_lower:MODEL}}     the exact binomial 95 % lower bound for that
                                model's episode count, as a percentage

`{{clopper_lower:…}}` is the one figure here that is *computed* rather than
quoted, and it is computed from the document's own episode cell — so if 2/2
becomes 1/2 the interval on the slide moves with it instead of going stale in the
most flattering possible direction.

Three checks enforce it, and they fail on different edits:

  1. `literal_metric_problems` — on the SOURCE markdown, before anything is
     resolved. Any two-or-more-decimal value under 1 anywhere in the slide
     document, notes included, is refused with the placeholder that replaces it.
     This is the check that stops the obvious break: pasting `0.933` in.
  2. `eval_side_problems` — a figure for any model other than the one
     `docs/architecture.md` §3 says the engine loads must be marked
     evaluation-side IN THE SAME BLOCK, not merely somewhere on the slide. This
     is the check that stops the *dangerous* break: swapping
     `{{measured:xgboost:AUROC}}` for `{{measured:fused:AUROC}}`, which check 1
     cannot see because the value is still resolved from the document and is
     still a real published number. The block scope is load-bearing and was
     measured, not assumed: slide 2's band already says "measured
     evaluation-side" about the fused row, so the obvious slide-scoped form of
     this check returns ZERO problems on that swap and ships it. The exception is
     a row §4 itself labels a baseline (`baseline_keys`), where the word
     "baseline" beside the figure is disclosure enough.
  3. `stale_number_problems` — on the SAVED .pptx, reopened. Every such value on
     a rendered slide must still be a figure the document measures, or a value
     derivable from it. This is what makes `--verify-only` meaningful on a
     committed binary nobody rebuilt.

WHAT ELSE THIS REFUSES TO BUILD
-------------------------------
  * more than `IDEA_MAX_SLIDES` slides, or a section list that is not `SECTIONS`;
  * a required disclosure missing from the slide that carries the claim it
    qualifies (`REQUIRED_ON_SLIDE`) — pinned per slide, not deck-wide, for the
    reason `build_deck.py` documents at length: a disclosure two slides from its
    claim is a deleted disclosure;
  * any on-slide run matching a banned claim. `build_deck.BANNED_ON_SLIDE` is
    imported whole (forward forecasting is not validated; the fused headline is
    not the engine) and `IDEA_BANNED` adds the three this surface invites: that
    the engine is a cascade or ensemble, accuracy theatre, and calling the
    tamper-evident ledger a blockchain;
  * a slide that overflows its own height or passes the on-slide word cap;
  * any run below the type-size floors a projected slide needs.

Determinism, offline behaviour and the zip-timestamp pinning are `build_deck`'s,
imported: the same markdown produces a byte-identical .pptx, and nothing here
opens a socket.
"""

from __future__ import annotations

import argparse
import hashlib
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_deck as bd  # noqa: E402  (the shared machinery; not modified by this file)

# ---------------------------------------------------------------- the format

IDEA_MAX_SLIDES = 6

# The official SIH IDEA presentation format, in order. Read off the official
# SIH2024_IDEA_Presentation_Format.pdf and cross-checked against the SIH2026
# revision; `docs/idea_round.md` carries the pointer list and the provenance.
SECTIONS: tuple[str, ...] = (
    "Title Page",
    "Proposed Solution",
    "Technical Approach",
    "Feasibility and Viability",
    "Impact and Benefits",
    "Research and References",
)

DECK_TITLE = "SIH26153 — idea round — AI based Network Attack Forecasting from Network Traffic Data"
FOOTER_LEFT = ("SIH26153 · NTRO · idea round · built from docs/idea_round.md by "
               "scripts/build_idea_deck.py")
SOURCE_DIGEST_KEY = "sih26-idea-deck-source-sha256"

# ------------------------------------------------- disclosures, pinned by slide
#
# Same rule `build_deck.py` learned the hard way: a disclosure is pinned to the
# slide that carries the claim it qualifies, never to the deck. The lead-time
# qualification belongs on slide 5 because slide 5 is where the number is
# printed; moving it to slide 4 is the same as deleting it.
#
# Phrases are compared through `build_deck.normalise`, so they match the source
# markdown and the runs read back out of the saved package alike.

REQUIRED_ON_SLIDE: dict[int, tuple[tuple[str, str], ...]] = {
    1: (
        ("sih26153", "the problem statement ID a screening judge scores against"),
    ),
    2: (
        ("day-wise",
         "the split rule is what makes the headline a cross-family transfer "
         "result rather than a memorised episode. Without it the number on this "
         "slide is just a number"),
        ("identical feature matrix",
         "the logistic baseline is only a fair comparison because it is trained "
         "and scored on the same matrix. Quoting the gap without saying so is the "
         "flattering version"),
    ),
    3: (
        ("one booster",
         "the engine is a SINGLE XGBoost model, not a cascade. This slide draws "
         "two lanes and this phrase is what stops a reader inferring that both "
         "of them run"),
    ),
    4: (
        ("n = 2 attack episodes",
         "the whole lead-time result rests on two episodes of one attacker host. "
         "The risk table is where a screening judge is told so"),
        ("3 of 7 kill-chain stages",
         "recon, lateral_movement and exfiltration have zero labelled windows; a "
         "deck that implies a validated seven-class classifier has overclaimed"),
    ),
    5: (
        ("horizon-0",
         "the lead time comes from the horizon-0 classifier — precursor detection "
         "on live windows — not from forecasting k steps ahead"),
        ("one attacker host",
         "n = 2, both sessions of one host. This is the slide that prints the "
         "lead time, so this is the slide that has to carry it"),
        ("matched-budget random baseline",
         "at this alert budget the lead time does not separate from a "
         "matched-budget random baseline. Printing the lead time without that is "
         "the largest overclaim this deck could make"),
    ),
    6: (
        ("cse-cic-ids2018",
         "the dataset the entire evaluation rests on has to be cited by name"),
    ),
}

# ------------------------------------------------- claims this surface invites

_ENGINE = r"\b(?:engine|production|demo|deployed\s+(?:model|scorer|system))\b"

IDEA_BANNED: tuple[bd.BannedOnSlide, ...] = (
    bd.BannedOnSlide(
        slug="engine-is-a-cascade",
        why=(
            "The deployed engine is a SINGLE XGBoost model -- one booster, one "
            "predict_proba call -- with two variants selected by input format "
            "(PCAP: all 30 features, CSV: flow-only). That is a variant choice, "
            "not a chain. Slide 3 draws two lanes, and 'cascade' or 'ensemble' is "
            "exactly the word a reader reaches for when they see two lanes, so "
            "the deck must never supply it."
        ),
        patterns=bd._rx(
            r"\bcascad(?:e|es|ed|ing)\b",
            r"\bensemble\b[^.]{0,60}" + _ENGINE,
            _ENGINE + r"[^.]{0,60}\bensemble\b",
            r"\bmulti[- ]?stage\s+(?:model|classifier|scorer|pipeline)\b",
            r"\btwo\s+models?\b[^.]{0,60}" + _ENGINE,
        ),
        retired=(
            "a cascade of two models",
            "the engine runs an ensemble of XGBoost and TGN",
            "an ensemble scores every host-window in production",
            "a multi-stage classifier scores each window",
            "two models run in the engine",
        ),
    ),
    bd.BannedOnSlide(
        slug="accuracy-theatre",
        why=(
            "The idea round is judged against decks claiming 99 % accuracy with "
            "no caveats, and the temptation is to answer in kind. Every phrasing "
            "below is either a number this project has not measured, a "
            "head-to-head it has not run (docs/competitive.md §4: Suricata and "
            "Zeek have never been executed over the test day), or a capability it "
            "does not have -- nothing here blocks, prevents or remediates "
            "anything; it forecasts, explains and records."
        ),
        patterns=bd._rx(
            r"\b9\d(?:\.\d+)?\s*%\s*(?:accura|detect|precision|recall|f1)",
            r"\b100\s*%\s*(?:accura|detect|precision|recall)",
            r"\bzero\s+false\s+positives?\b",
            r"\bstate[- ]of[- ]the[- ]art\b",
            r"\boutperform(?:s|ed|ing)?\b",
            r"\bbeats?\b[^.]{0,40}\b(?:existing|commercial|literature|baselines?|"
            r"suricata|zeek|snort|ndr)\b",
            r"\breal[- ]time\s+(?:blocking|prevention|mitigation|response)\b",
            r"\b(?:auto(?:matic|mated)?|self)[- ]?(?:block|blocks|blocking|"
            r"remediat\w*|mitigat\w*)\b",
            r"\b(?:stops?|prevents?|blocks?)\s+(?:the\s+)?attacks?\b",
        ),
        retired=(
            "99.4 % detection accuracy",
            "100 % accuracy on the test day",
            "zero false positives in production",
            "a state-of-the-art network attack forecaster",
            "our model outperforms commercial NDR",
            "beats existing intrusion detection systems",
            "real-time blocking of malicious hosts",
            "automatically blocks the attacking host",
            "stops attacks before they complete",
        ),
    ),
    bd.BannedOnSlide(
        slug="ledger-is-a-blockchain",
        why=(
            "docs/limitations.md: 'The ledger is tamper-evident, not a "
            "distributed blockchain.' Hash chain + per-batch Merkle roots + "
            "anchored checkpoints, verifiable offline. A live chain push was "
            "deliberately scoped out (BUILD_PLAN M9) because the demo must run "
            "air-gapped. The theme is called Blockchain & Cybersecurity, which is "
            "precisely why this word is so easy to over-reach with -- naming the "
            "theme is fine, claiming the technology is not. 'Immutable' is banned "
            "with it: the chain is tamper-EVIDENT, and it has a documented "
            "tampering mode it cannot catch."
        ),
        patterns=bd._rx(
            # "the Blockchain & Cybersecurity" theme name is allowed through; a
            # claim to be running one is not.
            r"\b(?:our|the|a)\s+blockchain\b(?!\s*(?:&|and\b))",
            r"\bblockchain[- ]based\b",
            r"\bon\s+(?:the\s+)?(?:a\s+)?(?:public\s+)?blockchain\b",
            r"\bsmart\s+contracts?\b",
            r"\b(?:distributed|decentrali[sz]ed)\s+ledger\b",
            r"\bimmutable\b",
        ),
        retired=(
            "our blockchain records every forecast",
            "a blockchain-based audit trail",
            "every forecast is written on the blockchain",
            "a smart contract verifies each batch",
            "a distributed ledger of forecasts",
            "an immutable record of every forecast",
        ),
    ),
)

ALL_BANNED: tuple[bd.BannedOnSlide, ...] = bd.BANNED_ON_SLIDE + IDEA_BANNED

# --------------------------------------------------------------- placeholders

PLACEHOLDER = re.compile(r"\{\{\s*(?P<kind>[a-z_]+)\s*:\s*(?P<arg>[^}]+?)\s*\}\}")

# Two-or-more-decimal values under 1 are the shape every figure in the results
# table takes (AUROC, F1). Identifiers that merely LOOK like one are removed
# first, because an arXiv id and a DOI are references, not measurements, and a
# check that cries wolf on the references slide gets switched off.
IDENTIFIER = re.compile(
    r"(?:arxiv|doi|rfc|isbn)\s*:?\s*\S+"          # arXiv:1603.02754, doi:10.1109/…
    r"|\b10\.\d{4,}/\S+"                          # a bare DOI
    r"|\bv?\d+(?:\.\d+){2,}\b",                   # a dotted version, e.g. 1.63.0
    re.IGNORECASE,
)
LITERAL_METRIC = re.compile(r"(?<![\w.])\d*\.\d{2,}")

# How a slide may say a figure is not the deployed engine's. Any one of these on
# the slide satisfies `eval_side_problems`.
EVAL_SIDE_MARKS: tuple[str, ...] = (
    "evaluation-side", "eval-side", "evaluation-only", "measured in eval",
    "not what the engine", "never loaded by the engine",
)


@dataclass(frozen=True)
class Resolved:
    """One placeholder, what it resolved to, which model it spoke for, and the
    block it sits in. The block matters: a disclosure on the same slide as a
    figure but in a different block is a disclosure a fast reader never joins to
    it — the slide-level version of the mistake `build_deck.py` fixed at the
    deck level. `block` is -1 for a placeholder in the slide title."""

    slide: int
    block: int
    kind: str
    model: str
    value: str


def baseline_keys(measured: dict[str, dict[str, str]]) -> set[str]:
    """Models `docs/architecture.md` §4 itself labels a baseline.

    Read, never declared. The graded logistic-regression row is a PS deliverable
    printed *as* a baseline and a reader cannot mistake it for the engine, so
    "baseline" beside it is disclosure enough. Every other non-deployed row --
    the fused headline above all -- needs the eval-side words. Taking the
    distinction from the document means a row that stops calling itself a
    baseline stops being treated as one.
    """
    return {key for key, row in measured.items()
            if "baseline" in bd.normalise(list(row.values())[0])}


def metric_tokens(text: str) -> list[str]:
    """Every value in `text` shaped like a published metric. See IDENTIFIER."""
    return LITERAL_METRIC.findall(IDENTIFIER.sub(" ", text))


def resolve_measured(measured: dict[str, dict[str, str]], model: str,
                     column: str, where: str) -> str:
    row = measured.get(model)
    if row is None:
        raise SystemExit(
            f"{where} asks for model {model!r}, which docs/architecture.md §4 does "
            f"not measure. It measures {sorted(measured)}. This builder never "
            "invents a row."
        )
    # `_column_index` matches a NORMALISED keyword against normalised headers, so
    # the placeholder may be written `AUROC` or `auroc` and mean the same column.
    index = bd._column_index(list(row), bd.normalise(column))
    if index is None:
        raise SystemExit(
            f"{where} asks for column {column!r} of the {model!r} row, and "
            f"docs/architecture.md §4 has no such column. It has {list(row)}."
        )
    return list(row.values())[index].strip()


def lead_minutes(measured: dict[str, dict[str, str]], model: str, where: str) -> str:
    """The model's median-lead cell, in whole minutes. Converted, never typed.

    The document states lead time in seconds and a slide reads better in minutes.
    Doing that conversion by hand is how a figure and its source drift apart, so
    it happens here, from the cell, on every build.
    """
    cell = resolve_measured(measured, model, "lead", where)
    seconds = bd._leading_number(cell)
    if seconds is None:
        raise SystemExit(
            f"{where} wants {model!r}'s lead in minutes and docs/architecture.md "
            f"§4's lead cell is {cell!r}, which carries no number to convert."
        )
    return str(int(round(float(seconds) / 60.0)))


def clopper_lower(measured: dict[str, dict[str, str]], model: str, where: str) -> str:
    """Exact binomial (Clopper-Pearson) 95 % lower bound, as a percentage.

    Computed from the document's own `episodes` cell rather than quoted, so the
    interval cannot outlive the count it describes: if 2/2 ever becomes 1/2 the
    number on the slide moves with it. `k/n` is parsed from the cell; the bound is
    `Beta(alpha/2; k, n-k+1)`, and 0 for k = 0.

    This exists because "2 of 2 episodes" is the most flattering possible way to
    state a result from a sample of two, and the interval is what two-for-two
    honestly buys.
    """
    cell = resolve_measured(measured, model, "episodes", where)
    match = re.search(r"(\d+)\s*[/of]+\s*(\d+)", bd.plain(cell))
    if match is None:
        raise SystemExit(
            f"{where} wants an exact binomial interval for {model!r} and "
            f"docs/architecture.md §4's episodes cell is {cell!r}, which this "
            "builder cannot read as k/n. Nothing is guessed."
        )
    k, n = int(match.group(1)), int(match.group(2))
    if n <= 0 or k > n:
        raise SystemExit(f"{where}: {cell!r} is not a usable episode count (k={k}, n={n})")
    if k == 0:
        return "0.0"
    try:
        from scipy.stats import beta
    except ModuleNotFoundError:                      # pragma: no cover - scipy is pinned
        raise SystemExit(
            "scipy is needed to compute the exact binomial interval and is not "
            "importable. It is pinned in requirements.txt."
        )
    return f"{float(beta.ppf(0.025, k, n - k + 1)) * 100.0:.1f}"


def derivable_values(measured: dict[str, dict[str, str]]) -> set[str]:
    """Every metric-shaped value the DOCUMENT justifies appearing on a slide.

    The saved-file check compares against this rather than against what this run
    happened to substitute, so `--verify-only` means something on a .pptx this
    process did not build.
    """
    allowed: set[str] = set()
    for model, row in measured.items():
        for cell in row.values():
            allowed.update(metric_tokens(cell))
        where = f"the {model!r} row"
        allowed.update(metric_tokens(lead_minutes(measured, model, where)))
        allowed.update(metric_tokens(clopper_lower(measured, model, where)))
    return allowed


def substitute(slides: list[bd.Slide], measured: dict[str, dict[str, str]]
               ) -> list[Resolved]:
    """Resolve every placeholder in place. Returns what was resolved, per slide."""
    resolved: list[Resolved] = []

    for slide in slides:
        where = f"docs/idea_round.md slide {slide.number}"
        block_index = -1

        def replace(match: re.Match[str]) -> str:
            kind, arg = match.group("kind"), match.group("arg").strip()
            if kind == "measured":
                model, _, column = arg.partition(":")
                model, column = model.strip(), column.strip()
                if not column:
                    raise SystemExit(
                        f"{where}: {{{{measured:…}}}} needs MODEL:COLUMN, got {arg!r}"
                    )
                value = resolve_measured(measured, model, column, where)
            elif kind == "lead_min":
                model, value = arg, lead_minutes(measured, arg, where)
            elif kind == "clopper_lower":
                model, value = arg, clopper_lower(measured, arg, where)
            else:
                raise SystemExit(
                    f"{where}: unknown placeholder {{{{{kind}:…}}}}. This builder "
                    "knows measured, lead_min and clopper_lower — and it will not "
                    "pass an unknown one through as literal text."
                )
            resolved.append(Resolved(slide.number, block_index, kind, model, value))
            return value

        slide.title = PLACEHOLDER.sub(replace, slide.title)
        rebuilt: list[bd.Block] = []
        for block_index, block in enumerate(slide.blocks):
            rebuilt.append(bd.Block(
                block.kind,
                PLACEHOLDER.sub(replace, block.label),
                tuple(PLACEHOLDER.sub(replace, line) for line in block.lines),
            ))
        slide.blocks = rebuilt
    return resolved


# -------------------------------------------------------- source-level checks


def section_problems(slides: list[bd.Slide]) -> list[str]:
    """The deck is the official section list, in order, and nothing else."""
    problems: list[str] = []
    if len(slides) != len(SECTIONS):
        problems.append(
            f"the official SIH IDEA format is {len(SECTIONS)} slides and this "
            f"document has {len(slides)}. The sections are {list(SECTIONS)} — a "
            "screening judge scores against that list, so a seventh slide is not "
            "a bonus and a missing one is lost marks."
        )
    for index, slide in enumerate(slides):
        if index >= len(SECTIONS):
            problems.append(f"slide {slide.number} ({slide.title!r}) is past the "
                            f"{len(SECTIONS)}-slide official format")
            continue
        section = SECTIONS[index]
        if bd.normalise(section) not in bd.normalise(slide.title):
            problems.append(
                f"slide {slide.number} is titled {slide.title!r} and the official "
                f"section in position {index + 1} is {section!r}. Put the official "
                "section name in the title — the judge's score sheet uses it."
            )
        if slide.number != index + 1:
            problems.append(
                f"slide numbered {slide.number} sits in position {index + 1}; the "
                "official sections are ordered and this deck must be too."
            )
    return problems


def literal_metric_problems(source: str) -> list[str]:
    """No published figure may be typed into `docs/idea_round.md`, notes included.

    Notes are checked too, and deliberately. They are not projected, but they ARE
    part of a document a judge may be handed, and a stale number in the speaker
    notes is a stale number a presenter reads out loud.
    """
    problems: list[str] = []
    body = source.split("## Slide", 1)
    if len(body) < 2:
        return ["docs/idea_round.md has no '## Slide N — title' headings at all"]
    for number, line in enumerate(("## Slide" + body[1]).splitlines(), start=1):
        for token in metric_tokens(line):
            problems.append(
                f"line {number} of the slide body types the figure {token!r}. "
                "Every published figure is resolved from docs/architecture.md at "
                "build time — write {{measured:MODEL:COLUMN}}, {{lead_min:MODEL}} "
                "or {{clopper_lower:MODEL}} instead. The results table is being "
                "regenerated under a new anonymisation key, so a value typed here "
                "today is a value that ships stale in September.\n"
                f"      THE LINE: {line.strip()[:110]!r}"
            )
    return problems


# ------------------------------------------------- checks over rendered text
#
# These run twice: on the source (so a break is reported before anything is
# painted) and on the runs read back out of the saved .pptx (so the check is
# against what a judge opens, not against this builder's model of it).


def required_problems(per_slide: dict[int, list[str]], where: str) -> list[str]:
    problems: list[str] = []
    for number, needed in REQUIRED_ON_SLIDE.items():
        haystack = [bd.normalise(t) for t in per_slide.get(number, [])]
        joined = " ".join(haystack)
        for phrase, why in needed:
            if bd.normalise(phrase) not in joined:
                problems.append(
                    f"{where}: slide {number} does not carry {phrase!r}.\n"
                    f"      WHY IT IS PINNED TO SLIDE {number}: {why}."
                )
    return problems


def banned_problems(per_slide: dict[int, list[str]], where: str) -> list[str]:
    """Matched one run of text at a time, for `build_deck.py`'s stated reason:
    joining a slide's runs manufactures adjacencies that are not on the slide."""
    problems: list[str] = []
    for number, texts in sorted(per_slide.items()):
        for text in texts:
            flat = bd.normalise(text)
            for claim in ALL_BANNED:
                for pattern in claim.patterns:
                    hit = pattern.search(flat)
                    if hit:
                        problems.append(
                            f"{where}: slide {number} claims "
                            f"[{claim.slug}] {hit.group(0)!r}\n"
                            f"      IN: {flat[:130]!r}\n"
                            f"      WHY IT IS BANNED: {claim.why}"
                        )
    return problems


def eval_side_problems(resolved: list[Resolved], texts: dict, deployed_key: str,
                       baselines: set[str], where: str, by_block: bool) -> list[str]:
    """A figure for a model the engine does not run must be marked NEXT TO IT.

    This is the check that catches the break `literal_metric_problems` cannot
    see. Swapping `{{measured:xgboost:AUROC}}` for `{{measured:fused:AUROC}}`
    resolves cleanly, produces a real published number, and puts the eval-side
    headline exactly where a screening judge reads it as the thing they would
    run. No check on the *set* of numbers can see that, because every number is
    still one the document measures.

    `by_block` is the strong form and runs on the source: the disclosure must be
    in the SAME BLOCK as the figure. Slide 2 already carries "measured
    evaluation-side" in its band, so a slide-level check would wave the swap
    through -- the band is four blocks away from the headline and a fast reader
    never joins them. The saved-file pass falls back to the slide-level form,
    because a .pptx has runs and shapes, not blocks; it still catches a figure
    hand-moved to a slide that discloses nothing.
    """
    problems: list[str] = []
    scope = "block" if by_block else "slide"
    for item in resolved:
        if item.model == deployed_key:
            continue
        key = (item.slide, item.block) if by_block else item.slide
        text = bd.normalise(" ".join(texts.get(key, [])))
        marks = EVAL_SIDE_MARKS + (("baseline",) if item.model in baselines else ())
        if not any(mark in text for mark in marks):
            problems.append(
                f"{where}: slide {item.slide} prints {item.value!r} for model "
                f"{item.model!r}, and docs/architecture.md §3 says the engine runs "
                f"{deployed_key!r}. That {scope} says none of {list(marks)}, so a "
                "judge reads a figure the engine does not produce as the one their "
                "own demo run would. That is this project's single largest "
                "recurring overclaim, and the disclosure has to sit beside the "
                f"number, not elsewhere on the slide.\n      THE {scope.upper()}: "
                f"{text[:130]!r}"
            )
    return problems


def stale_number_problems(per_slide: dict[int, list[str]], allowed: set[str],
                          where: str) -> list[str]:
    """Every figure on a SAVED slide is still one the document justifies."""
    problems: list[str] = []
    for number, texts in sorted(per_slide.items()):
        for text in texts:
            for token in metric_tokens(text):
                if token.lstrip("0") not in {a.lstrip("0") for a in allowed}:
                    problems.append(
                        f"{where}: slide {number} shows the figure {token!r}, which "
                        "docs/architecture.md §4 does not measure and this builder "
                        f"cannot derive from it. Measured/derivable: {sorted(allowed)}\n"
                        f"      IN: {text[:110]!r}"
                    )
    return problems


# ------------------------------------------------------------------ rendering


class IdeaPainter(bd.Painter):
    """`build_deck.Painter`, with this deck's footer.

    The footer names the source and the builder so a judge holding a PDF can tell
    which of the two decks it is and what regenerates it. Overridden rather than
    parameterised because `build_deck.py` is not modified by this file.
    """

    def footer(self) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches

        self.textbox("deck-footer", bd.FOOTER_TOP_IN, 0.3, [FOOTER_LEFT],
                     bd.PT_FOOTER, bd.MUTED, width=bd.BODY_W_IN - 1.2)
        box = self.slide.shapes.add_textbox(
            Inches(bd.SLIDE_W_IN - bd.MARGIN_IN - 1.2), Inches(bd.FOOTER_TOP_IN),
            Inches(1.2), Inches(0.3),
        )
        box.name = "deck-footer-number"
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        self._runs(para, f"{self.number} / {self.total}", bd.PT_FOOTER, bd.MUTED)


def block_texts(slides: list[bd.Slide]) -> dict[tuple[int, int], list[str]]:
    """Text per (slide number, block index), for the block-scoped checks.

    Keyed on the index into `slide.blocks` — the same index `substitute` records
    on each `Resolved` — so notes blocks occupy indices too and the two sides
    cannot drift apart. The title is block -1.
    """
    out: dict[tuple[int, int], list[str]] = {}
    for slide in slides:
        out[(slide.number, -1)] = [slide.title]
        for index, block in enumerate(slide.blocks):
            out[(slide.number, index)] = [block.label, *block.lines]
    return out


def source_texts(slides: list[bd.Slide]) -> dict[int, list[str]]:
    """On-slide text per slide, from the source. Notes are excluded: nobody
    projects them, so they can neither satisfy a disclosure nor break one."""
    out: dict[int, list[str]] = {}
    for slide in slides:
        texts = [slide.title]
        for block in slide.on_slide:
            if block.label:
                texts.append(block.label)
            if block.kind in ("points", "chips"):
                texts.extend(bd.list_items(block.lines))
            elif block.kind == "table":
                texts.extend(cell for row in bd.table_rows(block.lines) for cell in row)
            elif block.kind in ("pipeline", "lane"):
                texts.extend(bd.pipeline_stages(block.lines))
            else:
                texts.append(" ".join(ln.strip() for ln in block.lines).strip())
        out[slide.number] = [t for t in texts if t.strip()]
    return out


def render(slides: list[bd.Slide], out: Path, source_digest: str) -> None:
    bd._import_pptx()
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    deck.slide_width = Inches(bd.SLIDE_W_IN)
    deck.slide_height = Inches(bd.SLIDE_H_IN)
    blank = deck.slide_layouts[6]
    layout_problems: list[str] = []

    for spec in slides:
        slide = deck.slides.add_slide(blank)
        painter = IdeaPainter(slide, spec.number, len(slides))
        painter.chrome()
        painter.title(spec.title)
        for block in spec.on_slide:
            text = " ".join(ln.strip() for ln in block.lines).strip()
            if block.kind == "kicker":
                painter.kicker(text)
            elif block.kind == "headline":
                painter.headline(text)
            elif block.kind == "points":
                painter.points(bd.list_items(block.lines),
                               ordered=bd.list_is_ordered(block.lines))
            elif block.kind == "chips":
                painter.chips(block.label, bd.list_items(block.lines))
            elif block.kind == "band":
                painter.band(block.label, text)
            elif block.kind == "pipeline":
                painter.pipeline(bd.pipeline_stages(block.lines))
            elif block.kind == "lane":
                style, _, caption = block.label.partition("|")
                painter.lane(style.strip().lower(), caption.strip(),
                             bd.pipeline_stages(block.lines))
            elif block.kind == "table":
                painter.table(bd.table_rows(block.lines), None)
            else:
                raise SystemExit(
                    f"docs/idea_round.md slide {spec.number} uses deck marker "
                    f"{block.kind!r}, which this builder does not paint."
                )
        painter.footer()

        if painter.y > bd.CONTENT_BOTTOM_IN:
            layout_problems.append(
                f"slide {spec.number} overflows: content reaches {painter.y:.2f}in "
                f"of {bd.CONTENT_BOTTOM_IN:.2f}in usable. Move prose into the "
                "unmarked paragraphs (they become the speaker notes) — do not "
                "shrink the type."
            )
        if painter.words > bd.MAX_ON_SLIDE_WORDS:
            layout_problems.append(
                f"slide {spec.number} carries {painter.words} on-slide words (cap "
                f"{bd.MAX_ON_SLIDE_WORDS}). A screening judge gives this deck about "
                "ninety seconds; a wall of text is not read at all."
            )
        notes = spec.notes
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        print(f"   slide {spec.number} [{SECTIONS[spec.number - 1]}]: "
              f"{painter.words} on-slide words, content to {painter.y:.2f}in, "
              f"{len(notes.split())} words of notes")

    if layout_problems:
        raise SystemExit(
            "the deck does not fit its own slides:\n  - " + "\n  - ".join(layout_problems)
        )

    core = deck.core_properties
    core.title = DECK_TITLE
    core.subject = "SIH 2026 idea round · problem statement SIH26153 (NTRO)"
    core.author = "SIH26153 team"
    core.last_modified_by = "scripts/build_idea_deck.py"
    core.comments = ("Generated from docs/idea_round.md — edit the markdown, not this "
                     "file. Export to PDF before uploading: the SIH portal takes PDF.")
    core.keywords = f"{SOURCE_DIGEST_KEY}={source_digest}"
    core.created = bd.FIXED_TIME
    core.modified = bd.FIXED_TIME
    core.revision = 1

    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(out))
    bd._pin_zip_timestamps(out)


# --------------------------------------------------------------- verification


def saved_texts(path: Path) -> tuple[dict[int, list[str]], list[str]]:
    """Runs read back out of the saved package, per slide, plus type-size faults."""
    from pptx import Presentation

    deck = Presentation(str(path))
    per_slide: dict[int, list[str]] = {}
    problems: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        texts: list[str] = []
        for name, run in bd._slide_runs(slide):
            if run.text.strip():
                texts.append(run.text)
            size = run.font.size.pt if run.font.size is not None else None
            if size is None:
                problems.append(f"slide {index}: a run in {name!r} has no explicit "
                                "size, so it inherits whatever PowerPoint decides")
            elif size < bd.MIN_ANY_PT:
                problems.append(f"slide {index}: {name!r} sets {size:g}pt, below the "
                                f"{bd.MIN_ANY_PT:g}pt floor")
            elif size < bd.MIN_CONTENT_PT and not name.startswith("deck-footer"):
                problems.append(f"slide {index}: content run in {name!r} sets "
                                f"{size:g}pt, below the {bd.MIN_CONTENT_PT:g}pt floor "
                                "a judge reads from the back of a room")
        if not texts:
            problems.append(f"slide {index} carries no text at all")
        per_slide[index] = texts
    return per_slide, problems


def verify(path: Path, measured: dict[str, dict[str, str]], deployed_key: str,
           resolved: list[Resolved], baselines: set[str]) -> tuple[int, int]:
    from pptx import Presentation

    count = len(Presentation(str(path)).slides)
    per_slide, problems = saved_texts(path)

    if count > IDEA_MAX_SLIDES:
        problems.append(f"{count} slides; the official SIH IDEA format caps it at "
                        f"{IDEA_MAX_SLIDES}")
    if count != len(SECTIONS):
        problems.append(f"{count} slides; the official section list has "
                        f"{len(SECTIONS)} entries and every one is scored")

    for index, section in enumerate(SECTIONS, start=1):
        joined = bd.normalise(" ".join(per_slide.get(index, [])))
        if bd.normalise(section) not in joined:
            problems.append(f"the saved slide {index} does not name the official "
                            f"section {section!r}")

    problems += required_problems(per_slide, "saved deck")
    problems += banned_problems(per_slide, "saved deck")
    problems += stale_number_problems(per_slide, derivable_values(measured),
                                      "saved deck")
    problems += eval_side_problems(resolved, per_slide, deployed_key, baselines,
                                   "saved deck", by_block=False)

    if problems:
        raise SystemExit(
            f"{path} does not pass its own checks:\n  - " + "\n  - ".join(problems)
        )
    return count, path.stat().st_size


def source_digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def stamped_digest(path: Path) -> str | None:
    from pptx import Presentation

    keywords = Presentation(str(path)).core_properties.keywords or ""
    match = re.search(rf"{SOURCE_DIGEST_KEY}=([0-9a-f]{{64}})", keywords)
    return match.group(1) if match else None


# ----------------------------------------------------------------------- main


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--source", type=Path, default=REPO / "docs" / "idea_round.md",
                        help="source markdown (default: docs/idea_round.md)")
    parser.add_argument("--out", type=Path,
                        default=REPO / "docs" / "deck" / "sih26153_idea_deck.pptx",
                        help="output .pptx (default: docs/deck/sih26153_idea_deck.pptx)")
    parser.add_argument("--architecture", type=Path, default=bd.ARCHITECTURE_DOC,
                        help="document every figure is resolved from — §3 for which "
                             "model the engine runs, §4 for the results table. "
                             "Default: docs/architecture.md")
    parser.add_argument("--verify-only", action="store_true",
                        help="check an EXISTING .pptx and build nothing")
    parser.add_argument("--dump-text", action="store_true",
                        help="print every run of text on every slide of the saved "
                             "deck, then exit 0. This is what a human checks the "
                             "deck's CONTENT against the facts with.")
    parser.add_argument("--expect-sha256", default=None, metavar="HEX",
                        help="fail unless the .pptx has this SHA-256")
    args = parser.parse_args(argv)

    # A Windows console is cp1252 and the deck's own bullet glyph (U+25AA) has no
    # mapping in it, so --dump-text died with a UnicodeEncodeError on slide 1 --
    # i.e. the one command whose whole job is to let a human read the deck could
    # not print the deck. Degrade the glyph, never the run.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")

    measured = bd.read_measured_rows(args.architecture)
    deployed_key = bd.read_deployed_row_key(args.architecture)
    baselines = baseline_keys(measured)
    print(f"{args.architecture} -> engine runs {deployed_key!r}; "
          f"measured rows {sorted(measured)}; rows the document calls a baseline "
          f"{sorted(baselines)}")

    if args.dump_text:
        if not args.out.exists():
            raise SystemExit(f"{args.out} does not exist; build it first")
        per_slide, _ = saved_texts(args.out)
        for index, texts in sorted(per_slide.items()):
            print(f"\n=== SLIDE {index} — official section: {SECTIONS[index - 1]!r} ===")
            for text in texts:
                print(f"  | {text}")
        return 0

    if args.verify_only:
        if not args.out.exists():
            raise SystemExit(f"{args.out} does not exist; drop --verify-only to build it")
        count, size = verify(args.out, measured, deployed_key, [], baselines)
        print(f"verified (no rebuild) {args.out}")
        print(f"   {count} slides (official cap {IDEA_MAX_SLIDES}), {size:,} bytes")
        _report_digest(args.out, args.expect_sha256)
        stamped = stamped_digest(args.out)
        current = source_digest(args.source) if args.source.exists() else None
        if stamped is None:
            print(f"   NOTE: no source stamp — nothing here can tell this from a "
                  "stale deck. Rebuild it.")
        elif current is not None and stamped != current:
            raise SystemExit(
                f"{args.out} is STALE: built from a {args.source.name} that hashed "
                f"{stamped}, and that file now hashes {current}. Every check above "
                "passed — a deck built from last week's markdown satisfies all of "
                "them and still shows a judge the slide the fix replaced."
            )
        else:
            print(f"   source stamp matches {args.source}")
        return 0

    if not args.source.exists():
        raise SystemExit(f"{args.source} does not exist")
    raw = args.source.read_text(encoding="utf-8")

    literal = literal_metric_problems(raw)
    if literal:
        raise SystemExit(
            f"{args.source} types a published figure:\n  - " + "\n  - ".join(literal)
        )

    slides = bd.parse_slides(raw)
    if not slides:
        raise SystemExit(
            f"{args.source} contains no '## Slide N — title' headings. This script "
            "never invents slide content."
        )
    sections = section_problems(slides)
    if sections:
        raise SystemExit(
            "this deck is not the official SIH IDEA format:\n  - "
            + "\n  - ".join(sections)
        )

    resolved = substitute(slides, measured)
    for item in resolved:
        print(f"   slide {item.slide}: {{{{{item.kind}:{item.model}…}}}} -> "
              f"{item.value!r}")

    texts = source_texts(slides)
    problems = (required_problems(texts, str(args.source))
                + banned_problems(texts, str(args.source))
                + eval_side_problems(resolved, block_texts(slides), deployed_key,
                                     baselines, str(args.source), by_block=True))
    if problems:
        raise SystemExit(
            f"{args.source} does not pass its own checks:\n  - "
            + "\n  - ".join(problems)
        )

    print(f"{args.source} -> {len(slides)} slides")
    render(slides, args.out, source_digest(args.source))
    count, size = verify(args.out, measured, deployed_key, resolved, baselines)
    print(f"-> {args.out}")
    print(f"   {count} slides (official cap {IDEA_MAX_SLIDES}), {size:,} bytes")
    print("   NOTE: the SIH portal takes PDF. Export this .pptx to PDF before upload.")
    _report_digest(args.out, args.expect_sha256)
    return 0


def _report_digest(path: Path, expected: str | None) -> None:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    print(f"   sha256 {digest}")
    if expected and digest != expected.strip().lower():
        raise SystemExit(
            f"{path} hashes {digest}, not the expected {expected.strip().lower()}."
        )


if __name__ == "__main__":
    sys.exit(main())
